from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/zj/Documents/AI 自媒体')
MD = ROOT / 'GPT-5.6介绍：从更强模型到工作型AI.md'
OUT = ROOT / 'GPT-5.6头条发布稿.docx'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_run_font(run, name='Arial Unicode MS', size=11, bold=False, color='000000'):
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn(f'w:{key}'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_text(p, text, bold=False, size=11, color='000000'):
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return r

def add_md_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    # Keep simple Markdown emphasis readable in Word.
    pos = 0
    for m in re.finditer(r'\*\*([^*]+)\*\*|`([^`]+)`', text):
        add_text(p, text[pos:m.start()])
        add_text(p, m.group(1) or m.group(2), bold=bool(m.group(1)))
        pos = m.end()
    add_text(p, text[pos:])
    return p

def add_image(doc, rel):
    path = (ROOT / rel).resolve()
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(str(path), width=Inches(6.0))

def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = 'Table Grid'
    for i, value in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, value.strip(), bold=True, size=9, color='FFFFFF')
        set_cell_shading(cell, '1F4E78')
        set_cell_margins(cell)
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, value.strip(), size=9)
            set_cell_margins(cells[i])
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = 1
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Arial Unicode MS'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Arial Unicode MS')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial Unicode MS')
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
    normal._element.rPr.rFonts.set(qn('w:cs'), 'Arial Unicode MS')
    normal.font.size = Pt(11)
    for name, size, color in [('Heading 1', 18, '000000'), ('Heading 2', 15, '000000'), ('Heading 3', 13, '434343')]:
        st = styles[name]
        st.font.name = 'Arial Unicode MS'
        st._element.rPr.rFonts.set(qn('w:ascii'), 'Arial Unicode MS')
        st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial Unicode MS')
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
        st._element.rPr.rFonts.set(qn('w:cs'), 'Arial Unicode MS')
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)

    lines = MD.read_text(encoding='utf-8').splitlines()
    i = 0
    skip_sources = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == '### 本文使用的来源':
            skip_sources = True
            i += 1
            continue
        if skip_sources:
            if line.startswith('## 第一部分：'):
                skip_sources = False
            else:
                i += 1
                continue
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(10)
            add_text(p, line[2:].strip(), bold=True, size=24)
        elif line.startswith('## '):
            p = doc.add_paragraph(line[3:].strip(), style='Heading 1')
            p.paragraph_format.page_break_before = False
        elif line.startswith('### '):
            doc.add_paragraph(line[4:].strip(), style='Heading 2')
        elif line.startswith('#### '):
            doc.add_paragraph(line[5:].strip(), style='Heading 3')
        elif line.startswith('!['):
            m = re.search(r'\]\(<\.\/([^>]+)>\)', line)
            if m:
                add_image(doc, m.group(1))
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cur = lines[i].strip()
                if not re.match(r'^\|\s*:?-+', cur):
                    rows.append([x.strip() for x in cur.strip('|').split('|')])
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        elif line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(7)
            add_text(p, line[2:], color='555555')
        elif re.match(r'^\d+\.\s+', line):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(4)
            add_text(p, re.sub(r'^\d+\.\s+', '', line))
        else:
            add_md_paragraph(doc, line)
        i += 1

    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    build()
